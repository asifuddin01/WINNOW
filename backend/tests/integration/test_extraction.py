"""Data extraction (guide 8.12): versioned forms, dual extraction, consensus and export."""

import csv
import io
import uuid
from pathlib import Path
from typing import Any

from fastapi import FastAPI
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    AuditLog,
    EntryStatus,
    ExtractionEntry,
    FullTextStatus,
    Record,
    TitleAbstractStatus,
)
from app.workers.exports import run_export
from tests.conftest import MemoryMailer
from tests.project_helpers import API, api, delete, get, patch, post
from tests.screening_helpers import settings, team

SCHEMA: dict[str, Any] = {
    "fields": [
        {
            "key": "design",
            "label": "Study design",
            "type": "select",
            "required": True,
            "options": ["RCT", "Cohort"],
        },
        {"key": "n", "label": "Participants", "type": "number", "integer": True, "unit": "people"},
        {
            "key": "arms",
            "label": "Arms",
            "type": "table",
            "columns": [
                {"key": "arm", "label": "Arm", "type": "short_text", "required": True},
                {"key": "mean", "label": "Mean", "type": "number"},
            ],
        },
    ]
}


async def included(db: AsyncSession, ids: list[uuid.UUID]) -> None:
    await db.execute(
        update(Record)
        .where(Record.id.in_(ids))
        .values(ta_final=TitleAbstractStatus.INCLUDED, ft_final=FullTextStatus.INCLUDED)
    )
    await db.commit()


async def published(client: AsyncClient, pid: str, *, dual: bool = False) -> dict[str, Any]:
    created = await post(
        client,
        f"/projects/{pid}/extraction-forms",
        {"name": "Trial data", "schema": SCHEMA, "dual": dual},
    )
    assert created.status_code == 201, created.text
    form = await post(client, f"/projects/{pid}/extraction-forms/{created.json()['id']}/publish")
    assert form.status_code == 200, form.text
    result: dict[str, Any] = form.json()
    return result


async def test_forms_are_built_published_locked_and_versioned(
    db_app: FastAPI, db: AsyncSession, mailer: MemoryMailer
) -> None:
    async with team(db_app, db, mailer, records=1) as t:
        base = f"/projects/{t.pid}/extraction-forms"
        # Only owners and admins build forms; a broken schema lists its problems.
        assert (await post(t.reviewer, base, {"name": "Mine", "schema": SCHEMA})).status_code == 403
        broken = await post(
            t.owner,
            base,
            {"name": "Broken", "schema": {"fields": [{"key": "X", "label": "", "type": "slider"}]}},
        )
        assert broken.status_code == 422
        assert broken.json()["code"] == "invalid_form"
        assert len(broken.json()["problems"]) == 3

        empty = await post(t.owner, base, {"name": "Trial data"})
        assert empty.status_code == 201
        draft = empty.json()
        assert (draft["version"], draft["published"], draft["family_id"]) == (1, False, draft["id"])
        early = await post(t.owner, f"{base}/{draft['id']}/publish")
        assert early.status_code == 422  # nothing to extract yet
        changed = await patch(t.owner, f"{base}/{draft['id']}", {"schema": SCHEMA, "dual": True})
        assert changed.status_code == 200, changed.text
        assert changed.json()["schema"]["fields"][1]["unit"] == "people"
        first = (await post(t.owner, f"{base}/{draft['id']}/publish")).json()
        assert first["published"] is True

        # Published is locked; a change is the next version, a draft copied from it.
        locked = await patch(t.owner, f"{base}/{first['id']}", {"name": "Renamed"})
        assert locked.status_code == 409
        second = await post(t.owner, f"{base}/{first['id']}/versions")
        assert second.status_code == 201
        assert (second.json()["version"], second.json()["family_id"]) == (2, first["id"])
        assert second.json()["dual"] is True
        again = await post(t.owner, f"{base}/{first['id']}/versions")
        assert again.status_code == 409  # version 2 is still a draft
        forms = (await get(t.reviewer, base)).json()
        assert [(f["version"], f["latest"]) for f in forms] == [(1, False), (2, True)]
        assert (await delete(t.owner, f"{base}/{first['id']}")).status_code == 409
        assert (await delete(t.owner, f"{base}/{second.json()['id']}")).status_code == 204
        actions = set(await db.scalars(select(AuditLog.action)))
        assert {
            "extraction.form_created",
            "extraction.form_published",
            "extraction.form_versioned",
        } <= actions


async def test_extracting_a_study_blinded_like_screening(
    db_app: FastAPI, db: AsyncSession, mailer: MemoryMailer
) -> None:
    async with team(db_app, db, mailer, records=2) as t:
        await settings(t.owner, t.pid, blind_mode=True)
        study, excluded = t.records
        draft = (
            await post(
                t.owner, f"/projects/{t.pid}/extraction-forms", {"name": "Draft", "schema": SCHEMA}
            )
        ).json()
        form = await published(t.owner, t.pid)
        base = f"/projects/{t.pid}/extraction-forms/{form['id']}"
        await included(db, [study])

        not_yet = await api(
            t.reviewer,
            "PUT",
            f"/projects/{t.pid}/extraction-forms/{draft['id']}/entries/{study}",
            {"data": {}},
        )
        assert not_yet.status_code == 409
        assert (
            await api(t.reviewer, "PUT", f"{base}/entries/{excluded}", {"data": {}})
        ).status_code == 409

        partial = await api(t.reviewer, "PUT", f"{base}/entries/{study}", {"data": {"n": "120"}})
        assert partial.status_code == 200, partial.text
        assert partial.json()["data"] == {"n": 120}
        refused = await api(
            t.reviewer,
            "PUT",
            f"{base}/entries/{study}",
            {"data": {"n": 1.5, "arms": [{"mean": 2}]}, "status": "submitted"},
        )
        assert refused.status_code == 422
        assert refused.json()["problems"] == {
            "design": "Required.",
            "n": "Should be a whole number.",
            "arms[1].arm": "Required.",
        }
        mine = {"design": "RCT", "n": 120, "arms": [{"arm": "Melatonin", "mean": 6.5}]}
        done = await api(
            t.reviewer, "PUT", f"{base}/entries/{study}", {"data": mine, "status": "submitted"}
        )
        assert done.json()["status"] == "submitted"
        await api(
            t.owner,
            "PUT",
            f"{base}/entries/{study}",
            {"data": {**mine, "n": 121}, "status": "submitted"},
        )

        # Blind mode: the reviewer sees their own only; the unblinded owner sees both.
        theirs = (await get(t.reviewer, f"{base}/entries/{study}")).json()
        assert [e["mine"] for e in theirs["entries"]] == [True]
        assert theirs["consensus"] is None
        everyone = (await get(t.owner, f"{base}/entries/{study}")).json()
        assert len(everyone["entries"]) == 2
        assert everyone["label"] == "Smith 2015"
        studies = (await get(t.reviewer, f"{base}/studies")).json()
        assert [(s["mine"], s["submitted"]) for s in studies] == [("submitted", None)]
        studies = (await get(t.owner, f"{base}/studies")).json()
        assert [(s["mine"], s["submitted"], s["consensus"]) for s in studies] == [
            ("submitted", 2, False)
        ]


async def test_two_extractions_become_one_consensus(
    db_app: FastAPI, db: AsyncSession, mailer: MemoryMailer
) -> None:
    async with team(db_app, db, mailer, records=1) as t:
        [study] = t.records
        await included(db, [study])
        form = await published(t.owner, t.pid, dual=True)
        base = f"/projects/{t.pid}/extraction-forms/{form['id']}"
        a = {
            "design": "RCT",
            "n": 120,
            "arms": [{"arm": "Melatonin", "mean": 6.5}, {"arm": "Placebo", "mean": 6}],
        }
        b = {
            "design": "RCT",
            "n": 118,
            "arms": [{"arm": "Melatonin", "mean": 6.5}, {"arm": "Placebo", "mean": 5.9}],
        }
        await api(t.owner, "PUT", f"{base}/entries/{study}", {"data": a, "status": "submitted"})
        await api(t.reviewer, "PUT", f"{base}/entries/{study}", {"data": b, "status": "submitted"})

        assert (await get(t.reviewer, f"{base}/consensus/{study}")).status_code == 403
        view = (await get(t.owner, f"{base}/consensus/{study}")).json()
        assert [(d["path"], d["values"]) for d in view["differences"]] == [
            ("n", [120, 118]),
            ("arms[2].mean", [6, 5.9]),
        ]
        assert view["agreed"] == {"design": "RCT"}
        assert view["consensus"] is None

        incomplete = await api(t.owner, "PUT", f"{base}/consensus/{study}", {"data": {"n": 119}})
        assert incomplete.status_code == 422
        agreed = {**a, "n": 119}
        saved = await api(t.owner, "PUT", f"{base}/consensus/{study}", {"data": agreed})
        assert saved.status_code == 200, saved.text
        assert saved.json()["data"]["n"] == 119
        statuses = set(await db.scalars(select(ExtractionEntry.status)))
        assert statuses == {EntryStatus.VERIFIED}
        record = (await get(t.owner, f"{base}/entries/{study}")).json()
        assert record["consensus"]["resolved_by"] == "Ada Lovelace"
        # A changed entry is no longer the one the consensus reconciled.
        await api(t.reviewer, "PUT", f"{base}/entries/{study}", {"data": b, "status": "submitted"})
        statuses = set(await db.scalars(select(ExtractionEntry.status)))
        assert statuses == {EntryStatus.VERIFIED, EntryStatus.SUBMITTED}


async def export_file(
    db_app: FastAPI, client: AsyncClient, pid: str, **body: Any
) -> tuple[dict[str, Any], bytes]:
    requested = await post(client, f"/projects/{pid}/exports", body)
    assert requested.status_code == 201, requested.text
    job = requested.json()
    await run_export(
        sessionmaker=db_app.state.sessionmaker,
        redis=db_app.state.redis,
        storage=db_app.state.storage,
        export_id=uuid.UUID(job["id"]),
    )
    finished = (await get(client, f"/projects/{pid}/exports/{job['id']}")).json()
    assert finished["status"] == "ready", finished
    response = await client.get(f"{API}/projects/{pid}/exports/{job['id']}/file")
    return finished, response.content


async def test_extracted_data_exports_long_or_wide(
    db_app: FastAPI, db: AsyncSession, mailer: MemoryMailer, tmp_path: Path
) -> None:
    async with team(db_app, db, mailer, records=3) as t:
        reconciled, single, twice = t.records
        await included(db, t.records)
        form = await published(t.owner, t.pid, dual=True)
        base = f"/projects/{t.pid}/extraction-forms/{form['id']}"
        entry = {"design": "RCT", "n": 10, "arms": [{"arm": "A", "mean": 1}]}
        for record in (reconciled, twice):
            await api(
                t.owner, "PUT", f"{base}/entries/{record}", {"data": entry, "status": "submitted"}
            )
            await api(
                t.reviewer,
                "PUT",
                f"{base}/entries/{record}",
                {"data": {**entry, "n": 11}, "status": "submitted"},
            )
        await api(
            t.owner,
            "PUT",
            f"{base}/entries/{single}",
            {"data": {"design": "Cohort"}, "status": "submitted"},
        )
        await api(t.owner, "PUT", f"{base}/consensus/{reconciled}", {"data": {**entry, "n": 12}})

        options = {"form_id": form["id"], "layout": "wide", "which": "final"}
        job, content = await export_file(
            db_app, t.owner, t.pid, kind="extraction", format="csv", extraction=options
        )
        assert job["filename"].endswith("-extraction-trial-data-v1-wide.csv")
        rows = list(csv.DictReader(io.StringIO(content.decode())))
        # The consensus, and the only extraction; the study extracted twice waits.
        assert {(r["record_id"], r["extractor"], r["n"]) for r in rows} == {
            (str(reconciled), "consensus", "12"),
            (str(single), "Ada Lovelace", ""),
        }
        assert "arms[1].mean" in rows[0]

        options = {"form_id": form["id"], "layout": "long", "which": "all"}
        job, content = await export_file(
            db_app, t.owner, t.pid, kind="extraction", format="xlsx", extraction=options
        )
        path = tmp_path / "long.xlsx"
        path.write_bytes(content)
        sheet = load_workbook(path).active
        assert sheet is not None
        header = [cell.value for cell in sheet[1]]
        assert header == [
            "record_id",
            "study",
            "extractor",
            "field",
            "label",
            "row",
            "column",
            "value",
            "unit",
        ]
        values = list(sheet.iter_rows(min_row=2, values_only=True))
        # Every submitted extraction and the consensus: two people's four values on two
        # studies, the consensus's four, and the single extraction's one.
        assert len(values) == 2 * 2 * 4 + 4 + 1
        assert sum(1 for row in values if row[2] == "consensus") == 4

        # A blinded reviewer's file holds their own extractions only.
        await settings(t.owner, t.pid, blind_mode=True)
        _, content = await export_file(
            db_app,
            t.reviewer,
            t.pid,
            kind="extraction",
            format="csv",
            extraction={**options, "layout": "wide"},
        )
        mine = list(csv.DictReader(io.StringIO(content.decode())))
        assert [r["n"] for r in mine] == ["11", "11"]  # the reviewer's own two
        assert "consensus" not in {r["extractor"] for r in mine}

        wrong = await post(
            t.owner,
            f"/projects/{t.pid}/exports",
            {"kind": "extraction", "format": "ris", "extraction": options},
        )
        assert wrong.status_code == 422
        missing = await post(
            t.owner, f"/projects/{t.pid}/exports", {"kind": "extraction", "format": "csv"}
        )
        assert missing.status_code == 422
        other = await post(
            t.owner,
            f"/projects/{t.pid}/exports",
            {
                "kind": "extraction",
                "format": "csv",
                "extraction": {**options, "form_id": str(uuid.uuid4())},
            },
        )
        assert other.status_code == 404
